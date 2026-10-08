"""Knowledge ingestion, BM25 retrieval and grounded AI answers on an isolated SQL database."""
import io
from datetime import UTC, datetime
from uuid import UUID

import httpx
import pytest
from fastapi.testclient import TestClient

from app.api.deps import get_current_staff
from app.core.config import settings
from app.main import app
from app.models.schema import AIResponse, Conversation, KnowledgeChunk, UserSession
from app.services import ai_service, knowledge_service, rag_service
from app.services.ai_service import AIService
from app.services.knowledge_service import ExtractionError, Segment, chunk_segments, extract_segments
from app.services.staff_auth_service import StaffIdentity
from tests.conftest import TEST_DEVICE_ID

URL = "/api/v1/knowledge"
HOURS = ("Giờ mở cửa thư viện\n\nThư viện mở cửa từ 7 giờ 00 đến 21 giờ 00, từ thứ Hai đến thứ Bảy. "
         "Chủ nhật thư viện nghỉ.")
RULES = ("Nội quy thư viện\n\nBạn đọc phải xuất trình thẻ sinh viên khi vào thư viện. "
         "Không mang đồ ăn vào phòng đọc. Giữ trật tự trong khu vực học nhóm.")


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out)); out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


@pytest.fixture(autouse=True)
def isolated_storage(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "media_storage_dir", tmp_path)
    monkeypatch.setattr(rag_service, "_cached", None)


@pytest.fixture
def staff(use_db):
    identity = StaffIdentity(UUID("0d0d0d0d-0000-4000-8000-0000000000bb"), "librarian", "Thủ thư", "librarian", 4_102_444_800)
    app.dependency_overrides[get_current_staff] = lambda: identity
    yield TestClient(app)
    app.dependency_overrides.pop(get_current_staff, None)


# --- extraction and chunking --------------------------------------------------------------------

def test_extracts_every_supported_format():
    import docx
    import openpyxl

    word = docx.Document(); word.add_paragraph("Phòng học nhóm ở tầng 2.")
    table = word.add_table(rows=1, cols=2); table.rows[0].cells[0].text = "WiFi"; table.rows[0].cells[1].text = "NLU-Library"
    word_bytes = io.BytesIO(); word.save(word_bytes)
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = "Lịch"
    sheet.append(["Ngày", "Giờ mở cửa"]); sheet.append(["Thứ Hai", "7:00"])
    book_bytes = io.BytesIO(); book.save(book_bytes)

    assert "tầng 2" in extract_segments(".docx", word_bytes.getvalue())[0].text
    assert "WiFi | NLU-Library" in extract_segments(".docx", word_bytes.getvalue())[0].text
    xlsx = extract_segments(".xlsx", book_bytes.getvalue())[0]
    assert xlsx.sheet_name == "Lịch" and "Giờ mở cửa: 7:00" in xlsx.text
    assert "Giờ mở cửa: 7:00" in extract_segments(".csv", "Ngày;Giờ mở cửa\nThứ Hai;7:00\n".encode())[0].text
    pdf = extract_segments(".pdf", minimal_pdf("Opening hours 7:00 to 21:00"))
    assert pdf[0].page_number == 1 and "Opening hours" in pdf[0].text


@pytest.mark.parametrize(("extension", "data", "message"), [
    (".txt", bytes([0x47, 0x69, 0xF2, 0x20, 0x6D, 0xE1]), "UTF-8"),  # cp1258 bytes, invalid UTF-8
    (".txt", b"   \n  ", "không có nội dung"),
    (".pdf", b"%PDF-1.4 broken", "PDF"),
    (".docx", b"PK\x03\x04broken", "Word"),
])
def test_extraction_failures_are_reported_in_vietnamese(extension, data, message):
    with pytest.raises(ExtractionError, match=message):
        extract_segments(extension, data)


def test_chunks_respect_size_overlap_and_page_boundaries():
    long_page = " ".join(f"Câu số {i} nói về quy định thư viện." for i in range(120))
    chunks = chunk_segments([Segment(long_page, page_number=1), Segment("Trang hai ngắn.", page_number=2)])
    assert len(chunks) > 3
    assert all(len(chunk.text) <= knowledge_service.CHUNK_TARGET_CHARS + knowledge_service.CHUNK_OVERLAP_CHARS + 2 for chunk in chunks)
    assert chunks[-1].text == "Trang hai ngắn." and chunks[-1].page_number == 2
    tail = chunks[0].text[-40:]
    assert tail in chunks[1].text  # overlap keeps context across the boundary


# --- retrieval ----------------------------------------------------------------------------------

def add(db, title, content):
    return knowledge_service.create_from_text(db, title=title, content=content)


def test_retrieval_finds_relevant_chunks_with_or_without_diacritics(sqlite_db):
    add(sqlite_db, "Giờ mở cửa", HOURS); add(sqlite_db, "Nội quy", RULES)
    for question in ("Thư viện mở cửa lúc mấy giờ?", "thu vien mo cua luc may gio"):
        results = rag_service.retrieve(sqlite_db, question)
        assert results and results[0].title == "Giờ mở cửa", question
    assert rag_service.retrieve(sqlite_db, "Có được mang đồ ăn vào không?")[0].title == "Nội quy"
    assert rag_service.retrieve(sqlite_db, "Hôm nay thời tiết thế nào?") == []


def test_retrieval_ignores_inactive_deleted_and_failed_documents(sqlite_db):
    hours = add(sqlite_db, "Giờ mở cửa", HOURS)
    assert rag_service.retrieve(sqlite_db, "giờ mở cửa")
    knowledge_service.update(sqlite_db, hours.id, is_active=False)
    assert rag_service.retrieve(sqlite_db, "giờ mở cửa") == []
    knowledge_service.update(sqlite_db, hours.id, is_active=True)
    assert rag_service.retrieve(sqlite_db, "giờ mở cửa")  # index rebuilt after the change
    knowledge_service.remove(sqlite_db, hours.id)
    assert rag_service.retrieve(sqlite_db, "giờ mở cửa") == []
    failed = knowledge_service.create_from_upload(sqlite_db, file_name="scan.pdf", data=minimal_pdf(""), mime_type="application/pdf")
    assert failed.processing_status == "failed" and "OCR" in failed.source.processing_error


# --- admin API ----------------------------------------------------------------------------------

def test_knowledge_api_requires_staff(use_db):
    client = TestClient(app)
    assert client.get(f"{URL}/documents").status_code == 401
    assert client.post(f"{URL}/documents/text", json={"title": "x", "content": "y"}).status_code == 401
    assert client.post(f"{URL}/search", json={"query": "giờ"}).status_code == 401


def test_staff_manage_documents_end_to_end(staff, use_db):
    upload = staff.post(f"{URL}/documents", data={"title": "Giờ mở cửa"},
                        files={"file": ("gio-mo-cua.txt", HOURS.encode(), "text/plain")})
    assert upload.status_code == 201, upload.json()
    document = upload.json()["data"]
    assert document["status"] == "processed" and document["chunk_count"] >= 1 and document["source_type"] == "TEXT"

    text = staff.post(f"{URL}/documents/text", json={"title": "Nội quy", "content": RULES}).json()["data"]
    listing = staff.get(f"{URL}/documents").json()["data"]
    assert listing["total"] == 2 and listing["max_upload_mb"] == settings.max_knowledge_upload_mb
    assert staff.get(f"{URL}/documents", params={"search": "nội"}).json()["data"]["total"] == 1

    detail = staff.get(f"{URL}/documents/{document['id']}").json()["data"]
    assert "21 giờ 00" in detail["chunks"][0]["text"]
    found = staff.post(f"{URL}/search", json={"query": "mấy giờ thư viện đóng cửa"}).json()["data"]
    assert found[0]["document_id"] == document["id"] and found[0]["index"] == 1

    renamed = staff.patch(f"{URL}/documents/{text['id']}", json={"title": "Nội quy 2026", "is_active": False}).json()["data"]
    assert renamed["title"] == "Nội quy 2026" and renamed["is_active"] is False
    assert staff.post(f"{URL}/documents/{document['id']}/reprocess").json()["data"]["status"] == "processed"
    assert staff.delete(f"{URL}/documents/{document['id']}").status_code == 200
    assert staff.get(f"{URL}/documents/{document['id']}").status_code == 404
    assert staff.get(f"{URL}/documents").json()["data"]["total"] == 1


@pytest.mark.parametrize(("name", "content", "status", "code"), [
    ("virus.exe", b"MZ", 415, "UNSUPPORTED_DOCUMENT_TYPE"),
    ("fake.pdf", b"not a pdf", 415, "DOCUMENT_CONTENT_MISMATCH"),
    ("empty.txt", b"", 413, "DOCUMENT_TOO_LARGE"),
])
def test_upload_validation(staff, name, content, status, code):
    response = staff.post(f"{URL}/documents", files={"file": (name, content, "application/octet-stream")})
    assert response.status_code == status and response.json()["error"]["code"] == code


def test_upload_size_limit_and_failed_processing_is_visible(staff, monkeypatch):
    monkeypatch.setattr(settings, "max_knowledge_upload_mb", 1)
    big = staff.post(f"{URL}/documents", files={"file": ("big.txt", b"a" * (1024 * 1024 + 1), "text/plain")})
    assert big.status_code == 413
    failed = staff.post(f"{URL}/documents", files={"file": ("scan.pdf", minimal_pdf(""), "application/pdf")})
    assert failed.status_code == 201
    assert failed.json()["data"]["status"] == "failed" and failed.json()["message"].startswith("Đã lưu tài liệu nhưng")


def test_text_document_validation(staff):
    assert staff.post(f"{URL}/documents/text", json={"title": " ", "content": "x"}).status_code == 422
    assert staff.post(f"{URL}/documents/text", json={"title": "x", "content": "y", "extra": 1}).status_code == 422
    assert staff.patch(f"{URL}/documents/{TEST_DEVICE_ID}", json={"title": "x"}).status_code == 404


# --- grounded answers ---------------------------------------------------------------------------

def test_mock_provider_answers_extractively_with_citation(sqlite_db, monkeypatch):
    monkeypatch.setattr(settings, "ai_provider", "mock")
    add(sqlite_db, "Giờ mở cửa", HOURS)
    context = rag_service.retrieve(sqlite_db, "Thư viện mở cửa lúc mấy giờ?")
    result = AIService().answer("Thư viện mở cửa lúc mấy giờ?", [], context)
    assert result.grounded and result.provider == "knowledge"
    assert "7 giờ 00 đến 21 giờ 00" in result.text and result.citations[0]["title"] == "Giờ mở cửa"
    no_context = AIService().answer("Thư viện mở cửa lúc mấy giờ?", [], [])
    assert not no_context.grounded and "chưa được cung cấp tài liệu chính thức" in no_context.text


class FakeGemini:
    def __init__(self, text): self.text, self.payload = text, None

    def __call__(self, url, headers, json, timeout):
        self.payload = json
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": self.text}]}}]},
                              request=httpx.Request("POST", url))


def test_gemini_receives_delimited_context_and_citations_are_parsed(sqlite_db, monkeypatch):
    monkeypatch.setattr(settings, "ai_provider", "gemini"); monkeypatch.setattr(settings, "gemini_api_key", "test")
    add(sqlite_db, "Giờ mở cửa", HOURS); add(sqlite_db, "Nội quy", RULES)
    context = rag_service.retrieve(sqlite_db, "Thư viện mở cửa mấy giờ, có cần thẻ sinh viên không?")
    fake = FakeGemini("Thư viện mở cửa từ 7 giờ đến 21 giờ [1]. Bạn cần mang thẻ sinh viên [2].")
    monkeypatch.setattr(ai_service.httpx, "post", fake)
    result = AIService().answer("Thư viện mở cửa mấy giờ, có cần thẻ sinh viên không?", [], context)
    prompt = fake.payload["contents"][-1]["parts"][0]["text"]
    assert prompt.startswith("<tai_lieu>\n[1] ") and "</tai_lieu>" in prompt and "Câu hỏi của khách:" in prompt
    assert result.grounded and "[1]" not in result.text and result.text.endswith("thẻ sinh viên.")
    assert [c["index"] for c in result.citations] == [1, 2]

    monkeypatch.setattr(ai_service.httpx, "post", FakeGemini(ai_service.UNAVAILABLE_ANSWER))
    refusal = AIService().answer("Mật khẩu WiFi là gì?", [], context)
    assert refusal.grounded is False and refusal.citations == []


def test_gemini_failure_with_context_falls_back_to_document_text(sqlite_db, monkeypatch):
    monkeypatch.setattr(settings, "ai_provider", "gemini"); monkeypatch.setattr(settings, "gemini_api_key", "test")
    add(sqlite_db, "Giờ mở cửa", HOURS)
    def broken(*args, **kwargs): raise httpx.ConnectError("offline")
    monkeypatch.setattr(ai_service.httpx, "post", broken)
    result = AIService().answer("giờ mở cửa", [], rag_service.retrieve(sqlite_db, "giờ mở cửa"))
    assert result.provider_error and result.grounded and "21 giờ 00" in result.text


def test_kiosk_ai_turn_persists_citations(use_db, monkeypatch):
    monkeypatch.setattr(settings, "ai_provider", "mock")
    db = use_db
    add(db, "Giờ mở cửa", HOURS)
    session = UserSession(device_id=TEST_DEVICE_ID, started_at=datetime.now(UTC), identified=False)
    db.add(session); db.flush()
    conversation = Conversation(session_id=session.id, started_at=datetime.now(UTC), status="active")
    db.add(conversation); db.commit()
    response = TestClient(app).post("/api/v1/ai/answer", json={
        "conversation_id": str(conversation.id), "session_id": str(session.id), "message_text": "Thư viện mở cửa lúc mấy giờ?"})
    assert response.status_code == 200, response.json()
    data = response.json()["data"]
    assert data["grounded"] is True and data["citations"][0]["title"] == "Giờ mở cửa"
    stored = db.query(AIResponse).one()
    assert stored.grounded and stored.citations[0]["chunk_id"] == str(db.query(KnowledgeChunk).first().id)
