"""Knowledge ingestion: validate uploads, extract text, chunk it, and track processing status.

Documents are processed synchronously inside the request: library documents are small, and a
failed extraction is reported on the document (`status="failed"`, `processing_error`) rather
than raised, so staff can see and retry it.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID, uuid4

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.errors import AppError
from app.models.schema import KnowledgeChunk, KnowledgeDocument, KnowledgeSource

# Extension → (source/document type label, expected leading bytes or None for text).
SUPPORTED_TYPES: dict[str, tuple[str, bytes | None]] = {
    ".pdf": ("PDF", b"%PDF"),
    ".docx": ("DOCX", b"PK\x03\x04"),
    ".xlsx": ("XLSX", b"PK\x03\x04"),
    ".txt": ("TEXT", None),
    ".md": ("MARKDOWN", None),
    ".csv": ("CSV", None),
}
CHUNK_TARGET_CHARS = 900
CHUNK_OVERLAP_CHARS = 150
MAX_CHUNKS_PER_DOCUMENT = 2000
MAX_TEXT_DOCUMENT_CHARS = 200_000
STORAGE_DIR = "knowledge"


class ExtractionError(ValueError):
    """The file was accepted but its text could not be extracted; shown to staff as-is."""


@dataclass(frozen=True)
class Segment:
    text: str
    page_number: int | None = None
    sheet_name: str | None = None


@dataclass(frozen=True)
class Chunk:
    text: str
    page_number: int | None
    sheet_name: str | None


# --- extraction -------------------------------------------------------------------------------

def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("Tệp văn bản phải dùng mã hóa UTF-8.") from exc


def _extract_pdf(data: bytes) -> list[Segment]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("PDF được đặt mật khẩu; vui lòng tải bản không khóa.")
        return [Segment(page.extract_text() or "", page_number=index)
                for index, page in enumerate(reader.pages, start=1)]
    except PdfReadError as exc:
        raise ExtractionError("Không đọc được tệp PDF (tệp hỏng hoặc không hợp lệ).") from exc


def _extract_docx(data: bytes) -> list[Segment]:
    from zipfile import BadZipFile

    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except (BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError("Không đọc được tệp Word (.docx).") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            blocks.append(" | ".join(cell for cell in cells if cell))
    return [Segment("\n\n".join(blocks))]


def _rows_to_text(rows: list[list[str]]) -> str:
    """Spreadsheet rows become "Header: value" lines so each chunk is understandable alone."""
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    if not body:
        return " | ".join(header)
    lines = []
    for row in body:
        pairs = [f"{header[i] if i < len(header) and header[i] else f'Cột {i + 1}'}: {value}"
                 for i, value in enumerate(row) if value]
        lines.append("; ".join(pairs))
    return "\n\n".join(lines)


def _extract_xlsx(data: bytes) -> list[Segment]:
    from zipfile import BadZipFile

    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ExtractionError("Không đọc được tệp Excel (.xlsx).") from exc
    try:
        return [Segment(_rows_to_text([["" if value is None else str(value).strip() for value in row]
                                       for row in sheet.iter_rows(values_only=True)]), sheet_name=sheet.title)
                for sheet in workbook.worksheets]
    finally:
        workbook.close()


def _extract_csv(data: bytes) -> list[Segment]:
    text = _decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [Segment(_rows_to_text([[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]))]


def extract_segments(extension: str, data: bytes) -> list[Segment]:
    if extension == ".pdf":
        segments = _extract_pdf(data)
    elif extension == ".docx":
        segments = _extract_docx(data)
    elif extension == ".xlsx":
        segments = _extract_xlsx(data)
    elif extension == ".csv":
        segments = _extract_csv(data)
    else:
        segments = [Segment(_decode_text(data))]
    segments = [segment for segment in segments if segment.text.strip()]
    if not segments:
        if extension == ".pdf":
            raise ExtractionError("PDF không có lớp văn bản (có thể là bản scan); hệ thống chưa hỗ trợ OCR.")
        raise ExtractionError("Tài liệu không có nội dung văn bản.")
    return segments


# --- chunking ---------------------------------------------------------------------------------

_SENTENCE_END = re.compile(r"(?<=[.!?…])\s+")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pieces(paragraph: str) -> list[str]:
    """Split a paragraph that is too long into sentences, and over-long sentences by length."""
    if len(paragraph) <= CHUNK_TARGET_CHARS:
        return [paragraph]
    pieces = []
    for sentence in _SENTENCE_END.split(paragraph):
        while len(sentence) > CHUNK_TARGET_CHARS:
            cut = sentence.rfind(" ", 0, CHUNK_TARGET_CHARS)
            cut = cut if cut > CHUNK_TARGET_CHARS // 2 else CHUNK_TARGET_CHARS
            pieces.append(sentence[:cut].strip()); sentence = sentence[cut:].strip()
        if sentence:
            pieces.append(sentence)
    return pieces


def _overlap_tail(text: str) -> str:
    if len(text) <= CHUNK_OVERLAP_CHARS:
        return text
    tail = text[-CHUNK_OVERLAP_CHARS:]
    space = tail.find(" ")
    return tail[space + 1:] if 0 <= space < len(tail) - 1 else tail


def chunk_segments(segments: list[Segment]) -> list[Chunk]:
    """Greedy paragraph packing to ~CHUNK_TARGET_CHARS with a short overlap; never crosses pages/sheets."""
    chunks: list[Chunk] = []
    for segment in segments:
        current = ""
        for paragraph in _normalize(segment.text).split("\n\n"):
            for piece in _pieces(paragraph.strip()):
                if not piece:
                    continue
                if current and len(current) + len(piece) + 2 > CHUNK_TARGET_CHARS:
                    chunks.append(Chunk(current, segment.page_number, segment.sheet_name))
                    overlap = _overlap_tail(current)
                    current = f"{overlap}\n\n{piece}" if overlap else piece
                else:
                    current = f"{current}\n\n{piece}" if current else piece
        if current:
            chunks.append(Chunk(current, segment.page_number, segment.sheet_name))
    if len(chunks) > MAX_CHUNKS_PER_DOCUMENT:
        raise ExtractionError(f"Tài liệu quá dài (hơn {MAX_CHUNKS_PER_DOCUMENT} đoạn); vui lòng tách nhỏ.")
    return chunks


# --- persistence ------------------------------------------------------------------------------

def _storage_root() -> Path:
    root = Path(settings.media_storage_dir).resolve() / STORAGE_DIR
    root.mkdir(parents=True, exist_ok=True)
    return root


def _validate_upload(file_name: str, data: bytes) -> str:
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_TYPES:
        raise AppError(415, "UNSUPPORTED_DOCUMENT_TYPE",
                       "Chỉ hỗ trợ tệp PDF, Word (.docx), Excel (.xlsx), TXT, Markdown hoặc CSV.")
    max_bytes = settings.max_knowledge_upload_mb * 1024 * 1024
    if not data or len(data) > max_bytes:
        raise AppError(413, "DOCUMENT_TOO_LARGE", f"Tệp phải có dung lượng từ 1 byte đến {settings.max_knowledge_upload_mb} MB.")
    signature = SUPPORTED_TYPES[extension][1]
    if signature is not None and not data.startswith(signature):
        raise AppError(415, "DOCUMENT_CONTENT_MISMATCH", "Nội dung tệp không khớp với định dạng của tên tệp.")
    return extension


def _process(db: Session, document: KnowledgeDocument, extension: str, data: bytes) -> None:
    source = document.source
    db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.document_id == document.id))
    try:
        chunks = chunk_segments(extract_segments(extension, data))
    except ExtractionError as exc:
        source.status = document.processing_status = "failed"
        source.processing_error = str(exc)
        return
    for index, chunk in enumerate(chunks):
        metadata = {"source_type": source.source_type}
        db.add(KnowledgeChunk(document_id=document.id, chunk_index=index, chunk_text=chunk.text,
                              page_number=chunk.page_number, sheet_name=chunk.sheet_name, metadata_json=metadata))
    source.status = document.processing_status = "processed"
    source.processing_error = None


def _create(db: Session, *, title: str, file_name: str, extension: str, data: bytes, mime_type: str | None) -> KnowledgeDocument:
    source_type = SUPPORTED_TYPES[extension][0]
    stored = _storage_root() / f"{uuid4().hex}{extension}"
    stored.write_bytes(data)
    source = KnowledgeSource(source_name=title, source_type=source_type, original_file_name=file_name[:500],
                             file_mime_type=(mime_type or None), file_size=len(data),
                             storage_path=str(stored), status="processing")
    db.add(source); db.flush()
    document = KnowledgeDocument(source_id=source.id, source=source, title=title, document_type=source_type,
                                 language="vi", version="1", is_active=True, processing_status="processing")
    db.add(document); db.flush()
    _process(db, document, extension, data)
    db.commit(); db.refresh(document)
    return document


def _clean_title(title: str | None, fallback: str) -> str:
    cleaned = re.sub(r"\s+", " ", (title or "").strip())
    if not cleaned:
        cleaned = re.sub(r"[_-]+", " ", Path(fallback).stem).strip() or "Tài liệu"
    return cleaned[:500]


def create_from_upload(db: Session, *, file_name: str, data: bytes, mime_type: str | None,
                       title: str | None = None) -> KnowledgeDocument:
    file_name = Path(file_name or "tai-lieu").name
    extension = _validate_upload(file_name, data)
    return _create(db, title=_clean_title(title, file_name), file_name=file_name,
                   extension=extension, data=data, mime_type=mime_type)


def create_from_text(db: Session, *, title: str, content: str) -> KnowledgeDocument:
    content = content.strip()
    if not content:
        raise AppError(422, "EMPTY_DOCUMENT", "Nội dung tài liệu không được để trống.")
    if len(content) > MAX_TEXT_DOCUMENT_CHARS:
        raise AppError(413, "DOCUMENT_TOO_LARGE", f"Nội dung tối đa {MAX_TEXT_DOCUMENT_CHARS:,} ký tự.")
    cleaned = _clean_title(title, "tai-lieu")
    return _create(db, title=cleaned, file_name=f"{cleaned[:80]}.txt", extension=".txt",
                   data=content.encode("utf-8"), mime_type="text/plain")


def get_document(db: Session, document_id: UUID) -> KnowledgeDocument:
    document = db.get(KnowledgeDocument, document_id)
    if document is None or document.deleted_at is not None:
        raise AppError(404, "DOCUMENT_NOT_FOUND", "Không tìm thấy tài liệu.")
    return document


def reprocess(db: Session, document_id: UUID) -> KnowledgeDocument:
    document = get_document(db, document_id)
    source = document.source
    path = Path(source.storage_path or "")
    root = _storage_root()
    if not source.storage_path or root not in path.resolve().parents or not path.is_file():
        raise AppError(409, "DOCUMENT_FILE_MISSING", "Không còn tệp gốc để xử lý lại; vui lòng tải tài liệu lên lại.")
    _process(db, document, path.suffix.lower(), path.read_bytes())
    db.commit(); db.refresh(document)
    return document


def update(db: Session, document_id: UUID, *, title: str | None = None, is_active: bool | None = None) -> KnowledgeDocument:
    document = get_document(db, document_id)
    if title is not None:
        document.title = document.source.source_name = _clean_title(title, document.title)
    if is_active is not None:
        document.is_active = is_active
    db.commit(); db.refresh(document)
    return document


def remove(db: Session, document_id: UUID) -> None:
    """Soft delete: the chunks stay for citation history but are never retrieved again."""
    document = get_document(db, document_id)
    now = datetime.now(UTC)
    document.deleted_at = document.source.deleted_at = now
    document.is_active = False
    db.commit()


def chunk_counts(db: Session, document_ids: list[UUID]) -> dict[UUID, int]:
    if not document_ids:
        return {}
    rows = db.execute(select(KnowledgeChunk.document_id, func.count(KnowledgeChunk.id))
                      .where(KnowledgeChunk.document_id.in_(document_ids))
                      .group_by(KnowledgeChunk.document_id)).all()
    return {document_id: count for document_id, count in rows}


def list_documents(db: Session, *, search: str = "", status: str | None = None,
                   offset: int = 0, limit: int = 50) -> tuple[list[KnowledgeDocument], int]:
    query = select(KnowledgeDocument).where(KnowledgeDocument.deleted_at.is_(None))
    if search.strip():
        query = query.where(KnowledgeDocument.title.ilike(f"%{search.strip()}%"))
    if status:
        query = query.where(KnowledgeDocument.processing_status == status)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.scalars(query.order_by(KnowledgeDocument.created_at.desc()).offset(offset).limit(limit)).all()
    return list(rows), total
